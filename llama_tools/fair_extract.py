#!/usr/bin/env python3
"""Extract source-backed specimen/assay metadata and ask local Llama for proposals.

No ID mapping is delegated to the model. Source files are mounted read-only in
network-isolated bubblewrap workers; outputs always live in a new directory.
"""
import argparse
import codecs
from collections import defaultdict
import csv
from datetime import date, datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import re
import shutil
import sqlite3
import subprocess
import sys

sys.path.insert(0, str(Path(__file__).resolve().parent))
import rd_inventory as inventory
import sample_catalog as catalog_tools

DEFAULT_MODEL = Path(__file__).resolve().parent / 'models/Meta-Llama-3.1-8B-Instruct-Q5_K_M.gguf'
FIELDS = {'person_id_candidate', 'specimen_id', 'specimen_parent_id', 'specimen_label',
          'material', 'brain_region', 'laterality', 'tissue_block', 'aliquot_id',
          'collection_date', 'quantity', 'quantity_unit', 'assay_id', 'assay_type',
          'assay_platform', 'assay_status', 'processing_level', 'data_location',
          'reference_build', 'analysis_inclusion'}
ASSAYS = {'genotyping_array', 'wes', 'wgs', 'targeted_sequencing',
          'long_read_unspecified', 'other', 'unknown', None}
ROLES = {'brain_specimen', 'blood_specimen', 'genetic_assay', 'mixed', 'unknown'}
RELEVANT = re.compile(r'brain|region|hemisphere|laterality|material|specimen|barcode|aliquot|'
                      r'block|sample.?name|collection|quantity|volume|tissue|'
                      r'(?<![a-z])wes(?![a-z])|(?<![a-z])wgs(?![a-z])|gwas|exome|pacbio|genotyp|chip|imputation|'
                      r'data.?path|data.?director|reference|genome.?build|platform|assay|'
                      r'omics.?measured|dna|kit', re.I)
MATERIAL_BLOOD = {'blood', 'whole_blood', 'plasma', 'serum', 'buffy_coat', 'pbmc'}


def dump(path, value):
    with path.open('x', encoding='utf-8') as handle:
        json.dump(value, handle, indent=2, ensure_ascii=True)
        handle.write('\n')


def json_lines(path):
    with path.open(encoding='utf-8') as handle:
        for line in handle:
            if line.strip():
                yield json.loads(line)


def record_id(path, sheet):
    return hashlib.sha256(json.dumps([path, sheet], ensure_ascii=True).encode()).hexdigest()[:20]


def unsafecsv(value):
    """Reverse the existing catalogue writer's spreadsheet-formula protection."""
    return value[1:] if len(value) > 1 and value[0] == "'" and value[1] in '=+-@\t\r' else value


def make_plan(catalog, source_map=None, limit=None):
    selected = defaultdict(set)
    if source_map:
        with source_map.open(newline='', encoding='utf-8') as handle:
            for item in csv.DictReader(handle, delimiter='\t'):
                key = (unsafecsv(item['source_file']), unsafecsv(item['sheet']))
                selected[key].add(int(item['column_number']))
    sources = []
    seen = set()
    for schema in json_lines(catalog / 'schemas.jsonl'):
        if schema.get('kind') != 'schema' or schema.get('status') != 'selected':
            continue
        path, sheet = schema['path'], schema['sheet']
        catalog_tools.relative_path(path)
        key = (path, sheet)
        if key in seen or schema.get('role') == 'key':
            continue
        seen.add(key)
        if Path(path).suffix.lower() not in catalog_tools.EXTENSIONS:
            continue
        columns = schema['columns']
        semantic = sorted(selected[key]) if source_map else [
            i for i, header in enumerate(columns, 1) if RELEVANT.search(header)]
        if source_map and key not in selected:
            continue
        if not semantic:
            continue
        if any(i < 1 or i > len(columns) for i in semantic):
            raise ValueError(f'Source map column outside schema: {path} [{sheet}]')
        family = Path(path).parts[0].casefold()
        if not source_map and not any(word in family for word in ('brain', 'blood', 'biomaterial', 'genetic')):
            continue
        ids = schema.get('id_columns', [])
        id_indices = {i for item in ids for i in item['column_indices']}
        # Only semantic cells and explicitly recognized IDs are retained, not phenotypes.
        scope_indices = {columns.index(name) + 1 for item in ids for name in item.get('scope_columns', [])}
        read_columns = sorted(set(semantic) | id_indices | scope_indices)
        sources.append({'source_id': record_id(path, sheet), 'source_file': path,
                        'sheet': sheet, 'header_row': schema['header_row'],
                        'columns': columns, 'semantic_columns': semantic,
                        'read_columns': read_columns, 'id_columns': ids,
                        'requires_review': True})
    sources.sort(key=lambda item: (item['source_file'], item['sheet']))
    if limit:
        sources = sources[:limit]
    if not sources:
        raise ValueError('No selected specimen/assay tables found in schemas.jsonl')
    return {'schema_version': 1, 'catalog': str(catalog), 'requires_review': True,
            'selection': 'explicit source map' if source_map else 'relevant headers in brain/blood/biomaterial/genetics folders',
            'identity_policy': 'Existing confirmed row links, verified against current source ID cells only',
            'sources': sources}


def cell_value(cell):
    value = getattr(cell, 'value', cell)
    kind = getattr(cell, 'data_type', None)
    if kind in {'f', 'e', 'invalid_style'}:
        return None, 'formula_or_unreadable_cell'
    if value is None:
        return None, 'missing'
    if isinstance(value, (datetime, date)):
        return value.isoformat(), None
    if isinstance(value, (str, int, float, bool)):
        text = str(value)
        if len(text) > 4096:
            return text[:4096], 'value_truncated'
        return text, None
    return None, 'unsupported_cell_type'


def extract_rows(spec, rows):
    for number, cells in enumerate(rows, 1):
        if number < spec['header_row']:
            continue
        if number == spec['header_row']:
            headers = [str(getattr(cell, 'value', cell) or '').strip() for cell in cells]
            while headers and not headers[-1]:
                headers.pop()
            expected = list(spec['columns'])
            while expected and not expected[-1]:
                expected.pop()
            if headers != expected:
                raise ValueError('Source headers changed since catalogue; rebuild the catalogue before linking rows')
            inventory.emit({'kind': 'source_opened', 'source_id': spec['source_id']})
            break
    else:
        raise ValueError('Header row was not found')
    count = issues = 0
    for number, cells in enumerate(rows, spec['header_row'] + 1):
        values, flags = {}, {}
        for index in spec['read_columns']:
            cell = cells[index - 1] if index <= len(cells) else None
            value, flag = cell_value(cell)
            if value is not None:
                values[str(index)] = value
            if flag and flag != 'missing':
                flags[str(index)] = flag
        ids = []
        for item in spec['id_columns']:
            try:
                id_values = {i: catalog_tools.identifier(cells[i - 1], catalog_tools.MISSING)
                             for i in item['column_indices'] if i <= len(cells)}
                observed = set(id_values.values()) - {None}
                if len(observed) > 1:
                    raise ValueError('Conflicting duplicate ID cells')
                if observed:
                    # Preserve displayed leading zeroes and the catalogue's exact ID text.
                    for index, identifier in id_values.items():
                        if identifier is not None:
                            values[str(index)] = identifier
                    column = item['column']
                    if len(item['column_indices']) > 1:
                        column += ' [columns ' + ','.join(str(i) for i in item['column_indices']) + ']'
                    namespace = item['namespace']
                    scope = []
                    for name in item.get('scope_columns', []):
                        index = spec['columns'].index(name)
                        scope_value = catalog_tools.identifier(cells[index], catalog_tools.MISSING) if index < len(cells) else None
                        if scope_value is None:
                            raise ValueError('Missing namespace scope value')
                        scope.append([name, scope_value])
                    if scope:
                        namespace += '::scope=' + json.dumps(scope, ensure_ascii=True, separators=(',', ':'))
                    ids.append({'column': column, 'id': next(iter(observed)),
                                'namespace': namespace})
            except ValueError as exc:
                flags['id:' + item['column']] = str(exc)
                for index in item['column_indices']:
                    flags[str(index)] = 'invalid_identifier'
        if not values and not ids:
            continue
        count += 1
        issues += bool(flags)
        inventory.emit({'kind': 'row', 'source_id': spec['source_id'], 'row': number,
                        'values': values, 'ids': ids, 'flags': flags})
    inventory.emit({'kind': 'source_done', 'source_id': spec['source_id'],
                    'rows': count, 'rows_with_flags': issues})


def read_source(spec, config):
    """Read precisely one selected table, never traverse the source filesystem."""
    root, path = Path('/researchdrive'), spec['source_file']
    reason = inventory.content_skip_reason(path)
    if reason:
        raise ValueError('Source excluded by inventory policy: ' + reason)
    with os.fdopen(inventory.open_drive_path(root, path), 'rb') as handle:
        if Path(path).suffix.lower() == '.xlsx':
            import openpyxl
            if not openpyxl.DEFUSEDXML:
                raise RuntimeError('defusedxml is required for workbook parsing')
            # Formulas remain formulas and are excluded; do not evaluate cached values.
            book = openpyxl.load_workbook(handle, read_only=True, data_only=False, keep_links=False)
            try:
                if spec['sheet'] not in book.sheetnames:
                    raise ValueError('Worksheet no longer exists')
                sheet = book[spec['sheet']]
                sheet.reset_dimensions()
                extract_rows(spec, iter(sheet.iter_rows()))
            finally:
                book.close()
        else:
            rule = catalog_tools.file_rule(config, path, spec['sheet'])
            prefix = handle.read(4096)
            handle.seek(0)
            encoding = rule.get('encoding', 'utf-8-sig')
            if not rule.get('encoding'):
                if prefix.startswith((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE)):
                    encoding = 'utf-32'
                elif prefix.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
                    encoding = 'utf-16'
            if b'\x00' in prefix and encoding == 'utf-8-sig':
                raise ValueError('Binary text or unsupported encoding')
            delimiter = rule.get('delimiter')
            if delimiter is None:
                if Path(path).suffix.lower() == '.tsv':
                    delimiter = '\t'
                else:
                    try:
                        delimiter = csv.Sniffer().sniff(prefix.decode(encoding), delimiters=',;\t|').delimiter
                    except (csv.Error, UnicodeError):
                        delimiter = ','
            csv.field_size_limit(16 * 1024 * 1024)
            with io.TextIOWrapper(handle, encoding=encoding, newline='') as text:
                extract_rows(spec, iter(csv.reader(text, delimiter=delimiter)))


def sandbox(bwrap, mode, root=None, model=None, job=None, prompt=None, args=None):
    if mode not in {'read', 'model', 'preflight'}:
        raise ValueError('Only isolated source reading and model interpretation are supported')
    if mode in {'model', 'preflight'} and root is not None:
        raise ValueError('The FAIR Llama worker must never receive a raw-data mount')
    if mode == 'read' and root is None:
        raise ValueError('The source reader requires an explicitly read-only mount')
    command = inventory.sandbox_command(bwrap, 'probe', root=root, model=model,
                                         threads=args.threads if args else 4,
                                         gpu_layers=args.gpu_layers if args and mode in {'model', 'preflight'} else 0)
    index = command.index(str(Path(sys.executable).resolve()))
    binds = ['--ro-bind', str(Path(__file__).resolve()), '/app/fair_extract.py',
             '--ro-bind', str(Path(catalog_tools.__file__).resolve()), '/app/sample_catalog.py']
    if job:
        binds += ['--ro-bind', str(job), '/job.json']
    if prompt:
        binds += ['--ro-bind', str(prompt), '/prompt.txt']
    tail = [str(Path(sys.executable).resolve()), '-I', '-B', '/app/fair_extract.py', '_worker', mode]
    if mode in {'model', 'preflight'}:
        tail += [str(args.context), str(args.threads), str(args.gpu_layers), str(args.max_tokens)]
    return command[:index] + binds + tail


def validate_proposal(value, payload):
    required = {'schema_version', 'source_file', 'sheet', 'requires_review', 'bindings', 'value_mappings', 'interpretation_notes', 'unresolved_questions'}
    if not isinstance(value, dict) or set(value) != required or value.get('schema_version') != 1 or value.get('requires_review') is not True:
        raise ValueError('Invalid proposal envelope')
    if value.get('source_file') != payload['source_file'] or value.get('sheet') != payload['sheet']:
        raise ValueError('Model changed source identity')
    allowed = {item['column_number']: item['header'] for item in payload['columns']}
    bindings = value.get('bindings')
    if not isinstance(bindings, list) or len(bindings) > len(allowed) * 3:
        raise ValueError('Invalid or excessive bindings')
    seen_bindings = set()
    for item in bindings:
        if not isinstance(item, dict) or set(item) != {'column_number', 'header', 'target_field', 'assay_group', 'assay_type', 'confidence', 'evidence', 'requires_review'}:
            raise ValueError('Binding must be an object')
        index = item.get('column_number')
        if type(index) is not int or index not in allowed or item.get('header') != allowed[index]:
            raise ValueError('Model invented a column or changed its header')
        if item.get('target_field') not in FIELDS or item.get('confidence') not in {'low', 'medium', 'high'}:
            raise ValueError('Invalid field or confidence')
        if item.get('requires_review') is not True or not isinstance(item.get('evidence'), str) or not item['evidence']:
            raise ValueError('Binding lacks review flag or evidence')
        if item.get('assay_type') not in ASSAYS:
            raise ValueError('Invalid assay_type annotation')
        if item.get('assay_group') is not None and not isinstance(item['assay_group'], str):
            raise ValueError('Invalid assay group')
        key = (index, item['target_field'], item['assay_group'])
        if key in seen_bindings:
            raise ValueError('Duplicate or contradictory binding')
        seen_bindings.add(key)
    mappings = value.get('value_mappings')
    if not isinstance(mappings, list) or len(mappings) > 500:
        raise ValueError('Invalid or excessive value mappings')
    observed = defaultdict(set)
    for row in payload['preview_rows']:
        for index, raw in row['values'].items():
            observed[int(index)].add(raw)
    seen_mappings = set()
    for item in mappings:
        if not isinstance(item, dict) or set(item) != {'column_number', 'raw_value', 'normalized_value', 'evidence', 'requires_review'}:
            raise ValueError('Value mapping must be an object')
        index = item.get('column_number')
        if type(index) is not int or not isinstance(item.get('raw_value'), str) or item['raw_value'] not in observed[index]:
            raise ValueError('Value mapping was not observed in supplied preview')
        if not isinstance(item.get('normalized_value'), str) or not item['normalized_value']:
            raise ValueError('Invalid normalized term')
        if 'verified' in item['normalized_value'].casefold():
            raise ValueError('Llama cannot establish verified availability')
        key = (index, item['raw_value'])
        if key in seen_mappings:
            raise ValueError('Duplicate or contradictory value mapping')
        seen_mappings.add(key)
        if item.get('requires_review') is not True or not isinstance(item.get('evidence'), str) or not item['evidence']:
            raise ValueError('Value mapping lacks review flag or evidence')
    for field in ('interpretation_notes', 'unresolved_questions'):
        if not isinstance(value.get(field), list) or not all(isinstance(x, str) for x in value[field]):
            raise ValueError('Invalid proposal notes')
    return value


def worker():
    if sys.argv[2] == 'read':
        os.environ['OPENPYXL_LXML'] = 'False'
        with open('/job.json', encoding='utf-8') as handle:
            job = json.load(handle)
        try:
            read_source(job['source'], job['config'])
        except Exception as exc:
            inventory.emit({'kind': 'error', 'source_id': job['source']['source_id'], 'error': str(exc)})
        return
    import llama_cpp
    context, threads, gpu_layers, max_tokens = map(int, sys.argv[3:7])
    if gpu_layers and not llama_cpp.llama_supports_gpu_offload():
        raise RuntimeError('llama-cpp-python GPU backend is unavailable')
    if sys.argv[2] == 'preflight':
        inventory.emit({'preflight_ok': True, 'llama_cpp_version': getattr(llama_cpp, '__version__', 'unknown')})
        return
    from llama_cpp import Llama
    llm = Llama(model_path='/model.gguf', n_ctx=context, n_threads=threads,
                n_gpu_layers=gpu_layers, seed=0, verbose=gpu_layers != 0)
    prompt = Path('/prompt.txt').read_text(encoding='utf-8')
    for line in sys.stdin:
        payload = json.loads(line)
        try:
            # Never cut JSON or silently remove columns: trim only preview rows.
            while True:
                content = json.dumps(payload, ensure_ascii=True)
                used = len(llm.tokenize((prompt + content).encode('utf-8')))
                if used + max_tokens + 512 < context:
                    break
                if payload['preview_rows']:
                    payload['preview_rows'].pop()
                else:
                    raise ValueError('Headers and instructions exceed model context; increase --context')
            response = llm.create_chat_completion(
                messages=[{'role': 'system', 'content': prompt}, {'role': 'user', 'content': content}],
                temperature=0, max_tokens=max_tokens, response_format={'type': 'json_object'})
            choice = response['choices'][0]
            if choice.get('finish_reason') == 'length':
                raise ValueError('Model JSON truncated; increase --max-tokens and --context')
            value = validate_proposal(json.loads(choice['message']['content']), payload)
            inventory.emit({'source_id': payload['source_id'], 'proposal': value,
                            'preview_rows_used': len(payload['preview_rows'])})
        except (ValueError, TypeError, KeyError, RuntimeError) as exc:
            inventory.emit({'source_id': payload['source_id'], 'error': str(exc)})


def load_links(db, catalog, sources):
    db.execute('CREATE TABLE links(source TEXT,row INTEGER,column_name TEXT,identifier TEXT,namespace TEXT,person TEXT,status TEXT)')
    db.execute('CREATE INDEX link_rows ON links(source,row)')
    lookup = {(item['source_file'], item['sheet']): item['source_id'] for item in sources}
    for filename in ('id_matches.csv',):
        path = catalog / filename
        if not path.exists():
            continue
        with path.open(newline='', encoding='utf-8') as handle:
            for row in csv.DictReader(handle):
                row = {key: unsafecsv(value) for key, value in row.items()}
                source = lookup.get((row['path'], row['sheet']))
                if source:
                    db.execute('INSERT INTO links VALUES(?,?,?,?,?,?,?)',
                               (source, int(row['row']), row['column'], row['id'], row['namespace'], row['person_id'], row['status']))
    db.commit()


def row_identity(db, record):
    current = {(item['column'], item['id'], item['namespace']) for item in record['ids']}
    people, evidence, unresolved = set(), [], []
    for column, identifier, namespace, person, status in db.execute(
            'SELECT column_name,identifier,namespace,person,status FROM links WHERE source=? AND row=?',
            (record['source_id'], record['row'])):
        if (column, identifier, namespace) not in current:
            unresolved.append({'column': column, 'status': 'source_id_changed_or_missing'})
        elif status == 'confirmed' and person:
            people.add(person)
            evidence.append({'column': column, 'id': identifier, 'namespace': namespace, 'person_id': person})
        else:
            unresolved.append({'column': column, 'id': identifier, 'status': status})
    # Invalid IDs and cross-ID conflicts prevent a seemingly precise person link.
    invalid = any(key.startswith('id:') for key in record['flags'])
    stale = any(item['status'] == 'source_id_changed_or_missing' for item in unresolved)
    if stale:
        return None, 'stale_catalog_link', evidence, unresolved
    if len(people) == 1 and not invalid:
        return next(iter(people)), 'confirmed_existing_catalog_link', evidence, unresolved
    return None, 'conflicting_row_people' if len(people) > 1 else 'invalid_source_id' if invalid else 'unresolved', evidence, unresolved


def preview_identifier_header(header):
    """Withhold person/specimen labels and technical IDs from model examples."""
    return catalog_tools.id_header(header) or bool(re.search(
        r'(?:^|[^a-z])(?:id|name|donor|participant|person|subject|barcode|label|block|aliquot|path|directory|date|notes|comment|dna)(?:$|[^a-z])',
        header, re.I))


def assay_evidence(binding, raw, normalized):
    """An ID listed is weaker than completed data; negative/planned is never positive."""
    if binding.get('assay_type') in {None, 'unknown', 'other'}:
        return False
    field = binding['target_field']
    term = raw.strip().casefold()
    if term in catalog_tools.MISSING | {'0', '0.0', 'false', 'no', 'n', 'not_done', 'not done',
                                         'planned', 'pending', 'todo', 'to do', 'not_available'}:
        return False
    if normalized in {'not_done', 'reported_not_done', 'planned', 'unknown', 'not_available', 'negative', 'false'}:
        return False
    if field == 'assay_status':
        return normalized in {'reported_completed'}
    return field in {'assay_id', 'assay_type'}


def extract_fields(output, plan, proposals, db):
    sources = {item['source_id']: item for item in plan['sources']}
    people = defaultdict(lambda: {'brain_labels': set(), 'materials': set(), 'assays': set(),
                                 'brain_rows': set(), 'blood_rows': set(), 'genetic_rows': set(), 'assay_statuses': set(), 'files': set()})
    field_count = unresolved = 0
    with (output / 'extracted_fields.jsonl').open('x', encoding='utf-8') as fields, \
            (output / 'row_links.jsonl').open('x', encoding='utf-8') as links:
        for record in json_lines(output / 'observed_rows.jsonl'):
            spec = sources[record['source_id']]
            person, status, identity, issues = row_identity(db, record)
            links.write(json.dumps({'source_file': spec['source_file'], 'sheet': spec['sheet'],
                                    'row': record['row'], 'person_id': person, 'status': status,
                                    'evidence': identity, 'unresolved_links': issues,
                                    'source_flags': record['flags']}, ensure_ascii=True) + '\n')
            unresolved += person is None
            proposal = proposals.get(record['source_id'])
            if not proposal:
                continue
            mappings = {(item['column_number'], item['raw_value']): item['normalized_value']
                        for item in proposal['value_mappings']}
            for binding in proposal['bindings']:
                index = binding['column_number']
                raw = record['values'].get(str(index))
                if raw is None or raw.strip().casefold() in catalog_tools.MISSING:
                    continue
                normalized = mappings.get((index, raw))
                value = {'source_file': spec['source_file'], 'sheet': spec['sheet'],
                         'row': record['row'], 'column_number': index, 'header': binding['header'],
                         'person_id': person, 'person_link_status': status,
                         'target_field': binding['target_field'], 'raw_value': raw,
                         'proposed_normalized_value': normalized,
                         'assay_group': binding.get('assay_group'), 'proposed_assay_type': binding.get('assay_type'),
                         'confidence': binding['confidence'], 'requires_review': True,
                         'interpretation_evidence': binding['evidence'],
                         'cell_flag': record['flags'].get(str(index)),
                         'availability': 'metadata_observation_only_not_verified_data'}
                fields.write(json.dumps(value, ensure_ascii=True) + '\n')
                field_count += 1
                if person is None or record['flags'].get(str(index)):
                    continue
                item = people[person]
                item['files'].add(spec['source_file'])
                row_key = (record['source_id'], record['row'])
                field = binding['target_field']
                if field == 'brain_region':
                    item['brain_labels'].add(normalized or raw)
                    item['brain_rows'].add(row_key)
                if field == 'material':
                    item['materials'].add(normalized or raw)
                    material_term = re.sub(r'[ -]+', '_', normalized.casefold().strip()) if normalized else None
                    if material_term in MATERIAL_BLOOD:
                        item['blood_rows'].add(row_key)
                    if material_term in {'brain', 'brain_tissue'}:
                        item['brain_rows'].add(row_key)
                if field == 'assay_status':
                    item['assay_statuses'].add(json.dumps([binding.get('assay_group'), raw, normalized]))
                if assay_evidence(binding, raw, normalized):
                    item['assays'].add(binding['assay_type'])
                    item['genetic_rows'].add(row_key)
    with (output / 'person_overview.csv').open('x', newline='', encoding='utf-8') as handle:
        writer = csv.writer(handle)
        writer.writerow(['person_id', 'review_status', 'brain_region_labels_proposed',
                         'region_label_count_not_specimen_count', 'blood_materials_and_other_materials_proposed',
                         'genetic_assay_types_proposed', 'assay_status_evidence_raw_and_proposed', 'brain_evidence_rows', 'blood_evidence_rows',
                         'genetic_evidence_rows', 'brain_blood_genetics_candidate', 'source_files'])
        candidates = 0
        for person, item in sorted(people.items()):
            matched = bool(item['brain_rows'] and item['blood_rows'] and item['genetic_rows'])
            candidates += matched
            catalog_tools.safe_row(writer, [person, 'provisional_review_required', json.dumps(sorted(item['brain_labels'])),
                                           len(item['brain_labels']), json.dumps(sorted(item['materials'])),
                                           json.dumps(sorted(item['assays'])), json.dumps([json.loads(x) for x in sorted(item['assay_statuses'])]), len(item['brain_rows']),
                                           len(item['blood_rows']), len(item['genetic_rows']),
                                           'candidate_only' if matched else 'not_established',
                                           json.dumps(sorted(item['files']))])
    return {'extracted_fields': field_count, 'rows_without_confirmed_person': unresolved,
            'people_with_proposed_metadata': len(people), 'candidate_triple_matches': candidates}


def preflight_model(args, bwrap, model):
    """Check devices and inference dependencies before creating outputs or reading data."""
    command = sandbox(bwrap, 'preflight', model=model, args=args)
    result = subprocess.run(command, capture_output=True, text=True, timeout=60, env={})
    if result.returncode:
        raise RuntimeError('Model preflight failed before source reading: ' + result.stderr.strip())
    if not json.loads(result.stdout).get('preflight_ok'):
        raise RuntimeError('Model preflight did not confirm readiness')


def interpret_sources(args, bwrap, model, prompt, payloads, output, summary):
    proposals = {}
    print(f'Asking local Llama to interpret {len(payloads)} saved table schemas and bounded previews.', flush=True)
    if payloads:
        command = sandbox(bwrap, 'model', model=model, prompt=prompt, args=args)
        with (output / 'proposals.jsonl').open('x', encoding='utf-8') as handle:
            count = 0
            for result in inventory.child_lines(command, output / 'model.log', payloads=payloads):
                handle.write(json.dumps(result, ensure_ascii=True) + '\n')
                count += 1
                if 'proposal' in result:
                    proposals[result['source_id']] = result['proposal']
                    summary['sources_without_bindings'] += not result['proposal']['bindings']
                else:
                    summary['model_errors'] += 1
                print(f'Interpreted table {count}/{len(payloads)}.', flush=True)
            summary['model_errors'] += max(0, len(payloads) - count)
    return proposals


def resume_saved(args):
    """Interpret an unchanged extraction snapshot; never mount or re-read raw sources."""
    previous = args.resume_from.expanduser().resolve()
    output = args.output.expanduser().resolve()
    if getattr(args, 'root', None) is not None:
        raise ValueError('--resume-from uses saved reports only; omit --root')
    if output == previous or output.is_relative_to(previous):
        raise ValueError('Resume output must be a new directory outside the previous run')
    if output.exists():
        raise ValueError(f'Output already exists: {output}. Choose a new directory; previous results are preserved.')
    required = ['summary.json', 'source_plan.json', 'model_inputs.json', 'observed_rows.jsonl', 'links.sqlite']
    if any(not (previous / name).is_file() for name in required):
        raise ValueError('Previous run lacks saved extraction files; rerun extraction into a new output directory')
    old_summary = json.loads((previous / 'summary.json').read_text())
    plan = json.loads((previous / 'source_plan.json').read_text())
    payloads = json.loads((previous / 'model_inputs.json').read_text())
    if old_summary.get('read_errors') != 0 or len(payloads) != len(plan['sources']):
        raise ValueError('Saved extraction is incomplete or has read errors; rerun extraction into a new directory')
    if old_summary.get('source_write_access') is not False or old_summary.get('model_network_access') is not False:
        raise ValueError('Saved run does not record the required read-only isolation')
    expected = {item['source_id']: item for item in plan['sources']}
    if len(expected) != len(payloads) or {p.get('source_id') for p in payloads} != set(expected):
        raise ValueError('Saved table identities do not match the source plan')
    for payload in payloads:
        spec = expected[payload['source_id']]
        columns = [{'column_number': i, 'header': spec['columns'][i - 1]} for i in spec['semantic_columns']]
        if payload['source_file'] != spec['source_file'] or payload['sheet'] != spec['sheet'] or payload['columns'] != columns:
            raise ValueError('Saved payload no longer agrees with its source plan')
        ids = {i for item in spec['id_columns'] for i in item['column_indices']}
        visible = {i for i in spec['semantic_columns'] if i not in ids and not preview_identifier_header(spec['columns'][i - 1])}
        payload['preview_rows'] = [{'row': row['row'], 'values': {key: value for key, value in row['values'].items() if int(key) in visible}}
                                   for row in payload['preview_rows'][:args.preview_rows]]
    for suffix in ('-wal', '-journal'):
        sidecar = previous / ('links.sqlite' + suffix)
        if sidecar.exists() and sidecar.stat().st_size:
            raise ValueError('Saved link database has an active journal; wait for the previous run to finish')
    model, prompt = args.model.expanduser().resolve(), args.prompt.expanduser().resolve()
    if not model.is_file() or not prompt.is_file():
        raise ValueError('Model GGUF and prompt must exist')
    bwrap = shutil.which(args.bwrap)
    if not bwrap:
        raise RuntimeError('bubblewrap is required; no unsandboxed inference fallback')
    preflight_model(args, bwrap, model)
    inventory.check_sandbox(bwrap)
    output.mkdir(parents=True, mode=0o700, exist_ok=False)
    for name in ('source_plan.json', 'observed_rows.jsonl', 'links.sqlite'):
        shutil.copyfile(previous / name, output / name)
    dump(output / 'model_inputs.json', payloads)
    summary = dict(old_summary)
    summary.update({'started_utc': datetime.now(timezone.utc).isoformat(), 'complete': False,
                    'resumed_from': str(previous), 'source_tables_reread': False,
                    'snapshot_source_finished_utc': old_summary.get('finished_utc'),
                    'model': str(model), 'model_bytes': model.stat().st_size,
                    'model_errors': 0, 'sources_without_bindings': 0,
                    'prompt_sha256': hashlib.sha256(prompt.read_bytes()).hexdigest(),
                    'script_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
                    'python_version': sys.version,
                    'model_settings': {'context': args.context, 'seed': 0, 'temperature': 0,
                                       'threads': args.threads, 'gpu_layers': args.gpu_layers}})
    summary.pop('failure', None)
    db = sqlite3.connect(output / 'links.sqlite')
    try:
        summary['stage'] = 'model_interpretation'
        proposals = interpret_sources(args, bwrap, model, prompt, payloads, output, summary)
        summary['stage'] = 'overview'
        summary.update(extract_fields(output, plan, proposals, db))
        summary['complete'] = summary['model_errors'] == 0
        summary['stage'] = 'finished'
        print(f"Saved proposed matched-person overview: {output / 'person_overview.csv'}", flush=True)
        return 0 if summary['complete'] else 2
    except Exception as exc:
        summary['failure'] = str(exc)
        raise
    finally:
        db.close()
        summary['finished_utc'] = datetime.now(timezone.utc).isoformat()
        dump(output / 'summary.json', summary)


def run(args):
    os.umask(0o077)
    if getattr(args, 'resume_from', None) is not None:
        return resume_saved(args)
    catalog = args.catalog.expanduser().resolve()
    plan = make_plan(catalog, args.source_map, args.limit_sources)
    output = args.output.expanduser().resolve()
    if output.exists():
        raise ValueError(f'Output already exists: {output}. Choose a new directory or use --resume-from with a new --output.')
    root = args.root.expanduser().resolve() if args.root else None
    if output == catalog or output.is_relative_to(catalog):
        raise ValueError('Output must be a new directory outside the input catalogue')
    if root and (output == root or output.is_relative_to(root)):
        raise ValueError('Output must be outside ResearchDrive')
    if not args.plan_only:
        if root is None or not root.is_dir() or not args.allow_non_mount and not os.path.ismount(root):
            raise ValueError('Supply --root mounted ResearchDrive; --allow-non-mount is for deliberate fixtures only')
        for dependency in (Path(sys.prefix).resolve(), Path(__file__).resolve()):
            if dependency == root or dependency.is_relative_to(root):
                raise ValueError('Scripts and conda environment must be outside ResearchDrive')
        model = args.model.expanduser().resolve()
        prompt = args.prompt.expanduser().resolve()
        if not model.is_file() or not prompt.is_file():
            raise ValueError('Model GGUF and prompt must exist; models are never downloaded by this script')
        bwrap = shutil.which(args.bwrap)
        if not bwrap:
            raise RuntimeError('bubblewrap is required; no unsandboxed source-reading fallback')
        preflight_model(args, bwrap, model)
        inventory.check_sandbox(bwrap)
    output.mkdir(parents=True, mode=0o700, exist_ok=False)
    dump(output / 'source_plan.json', plan)
    if args.plan_only:
        print(f"Planned {len(plan['sources'])} tables: {output / 'source_plan.json'}; no source cells or model read.")
        return 0
    summary = {'started_utc': datetime.now(timezone.utc).isoformat(), 'complete': False,
               'requires_review': True, 'source_write_access': False, 'model_network_access': False,
               'catalog': str(catalog), 'root': str(root), 'model': str(model),
               'model_bytes': model.stat().st_size,
               'prompt_sha256': hashlib.sha256(prompt.read_bytes()).hexdigest(),
               'script_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
               'python_version': sys.version,
               'model_hash_note': 'Use the setup model provenance sidecar for full GGUF revision and SHA256; no multi-GB rehash here',
               'model_settings': {'context': args.context, 'seed': 0, 'temperature': 0,
                                  'threads': args.threads, 'gpu_layers': args.gpu_layers},
               'selected_sources': len(plan['sources']), 'read_errors': 0, 'model_errors': 0,
               'limit_sources': args.limit_sources, 'empty_previews': 0, 'sources_without_bindings': 0,
               'cohort_completeness': 'Not established; only selected readable tables, not all specimens or assays',
               'formula_cells': 'Excluded; formulas are never executed or treated as source ID evidence',
               'identifier_text': 'Catalogue normalization: trim surrounding whitespace, preserve case/punctuation and Excel displayed leading zeroes',
               'scope': 'All nonempty selected source rows; bounded Llama previews; no physical specimen or stored-data availability verification'}
    config_path = catalog / 'config_used.json'
    config = json.loads(config_path.read_text(encoding='utf-8')) if config_path.exists() else {}
    payloads = []
    db = sqlite3.connect(output / 'links.sqlite')
    try:
        load_links(db, catalog, plan['sources'])
        with (output / 'observed_rows.jsonl').open('x', encoding='utf-8') as rows, \
                (output / 'errors.jsonl').open('x', encoding='utf-8') as errors:
            for number, spec in enumerate(plan['sources'], 1):
                print(f"Reading table {number}/{len(plan['sources'])}: {spec['source_file']} [{spec['sheet']}]", flush=True)
                job = output / ('read-' + spec['source_id'] + '.json')
                dump(job, {'source': spec, 'config': config})
                preview, done = [], False
                id_indices = {i for item in spec['id_columns'] for i in item['column_indices']}
                # ID-like columns are withheld even when the old catalogue failed to recognize them.
                visible = [i for i in spec['semantic_columns'] if i not in id_indices
                           and not preview_identifier_header(spec['columns'][i - 1])]
                try:
                    command = sandbox(bwrap, 'read', root=root, job=job, args=args)
                    for record in inventory.child_lines(command, output / ('read-' + spec['source_id'] + '.log')):
                        if record['kind'] == 'row':
                            rows.write(json.dumps(record, ensure_ascii=True) + '\n')
                            if len(preview) < args.preview_rows:
                                values = {key: value for key, value in record['values'].items()
                                          if int(key) in visible and key not in record['flags']}
                                if values:
                                    preview.append({'row': record['row'], 'values': values})
                        elif record['kind'] == 'source_done':
                            done = True
                        elif record['kind'] == 'error':
                            errors.write(json.dumps(record, ensure_ascii=True) + '\n')
                            summary['read_errors'] += 1
                    if not done:
                        raise RuntimeError('Source extraction incomplete; records are partial and excluded from Llama interpretation')
                    summary['empty_previews'] += not preview
                    payloads.append({'source_id': spec['source_id'], 'source_file': spec['source_file'],
                                     'sheet': spec['sheet'], 'header_row': spec['header_row'],
                                     'columns': [{'column_number': i, 'header': spec['columns'][i - 1]}
                                                 for i in spec['semantic_columns']],
                                     'preview_rows': preview, 'identifier_values_withheld': True,
                                     'preview_scope': 'bounded examples; absent values are unknown'})
                except Exception as exc:
                    errors.write(json.dumps({'source_id': spec['source_id'], 'error': str(exc)}, ensure_ascii=True) + '\n')
                    summary['read_errors'] += 1
        # The model sandbox has no ResearchDrive mount and cannot request tools.
        dump(output / 'model_inputs.json', payloads)
        proposals = interpret_sources(args, bwrap, model, prompt, payloads, output, summary)
        summary.update(extract_fields(output, plan, proposals, db))
        summary['complete'] = summary['read_errors'] == 0 and summary['model_errors'] == 0
        print(f"Saved proposed matched-person overview: {output / 'person_overview.csv'}", flush=True)
        return 0 if summary['complete'] else 2
    finally:
        db.close()
        summary['finished_utc'] = datetime.now(timezone.utc).isoformat()
        dump(output / 'summary.json', summary)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    inputs = parser.add_mutually_exclusive_group(required=True)
    inputs.add_argument('--catalog', type=Path)
    inputs.add_argument('--resume-from', type=Path, help='Reuse a previous completed extraction snapshot in a new output; never read raw sources')
    parser.add_argument('--output', type=Path, required=True, help='New private output directory')
    parser.add_argument('--root', type=Path, help='Mounted ResearchDrive; unnecessary for --plan-only')
    parser.add_argument('--source-map', type=Path, help='Optional source_fields.tsv from the audit to restrict selected sources/columns')
    parser.add_argument('--plan-only', action='store_true', help='Select schemas only, without reading source cells or loading a model')
    parser.add_argument('--model', type=Path, default=DEFAULT_MODEL)
    parser.add_argument('--prompt', type=Path, default=Path(__file__).resolve().with_name('fair_extraction_prompt.txt'))
    parser.add_argument('--context', type=int, default=8192)
    parser.add_argument('--max-tokens', type=int, default=3072)
    parser.add_argument('--threads', type=int, default=4)
    parser.add_argument('--gpu-layers', type=int, default=0)
    parser.add_argument('--preview-rows', type=int, default=12)
    parser.add_argument('--limit-sources', type=int, help='Pilot subset; never a complete cohort report')
    parser.add_argument('--bwrap', default='bwrap')
    parser.add_argument('--allow-non-mount', action='store_true')
    args = parser.parse_args()
    if args.resume_from and (args.plan_only or args.source_map or args.limit_sources):
        parser.error('--resume-from reuses its saved source plan; omit --plan-only/--source-map/--limit-sources')
    if args.context < 4096 or args.max_tokens < 256 or args.max_tokens >= args.context - 1024:
        parser.error('Use context >=4096 and max-tokens between 256 and context-1024')
    if args.threads < 1 or not 1 <= args.preview_rows <= 100 or args.limit_sources is not None and args.limit_sources < 1:
        parser.error('threads/limit-sources must be positive; preview-rows must be 1..100')
    if args.gpu_layers < -1:
        parser.error('--gpu-layers must be -1 or nonnegative')
    return run(args)


if __name__ == '__main__':
    try:
        if len(sys.argv) > 1 and sys.argv[1] == '_worker':
            worker()
        else:
            sys.exit(main())
    except KeyboardInterrupt:
        sys.exit('Interrupted; partial output preserved')
    except Exception as exc:
        sys.exit(f'FAIR extraction failed: {exc}')
